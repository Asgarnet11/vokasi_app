"""
Utilitas ekspor data monitoring pelatihan vokasi ke format Excel (.xlsx).
Menggunakan openpyxl untuk menghasilkan laporan berformat profesional
lengkap dengan styling header, rekapitulasi per kejuruan, dan lembar data.
"""
import io
import pandas as pd
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def export_monitoring_excel(df: pd.DataFrame, kapasitas_default: int = 16) -> bytes:
    """Ekspor dataframe monitoring ke format Excel (.xlsx) dengan 2 sheet:
    1. Rekapitulasi: Ringkasan per Kejuruan (Bidang) & Capaian Kuota
    2. Data Pelatihan: Seluruh baris data pelatihan aktif/terfilter
    """
    wb = Workbook()

    # Warna tema (Navy / Slate profesional)
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    sub_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    sub_font = Font(name="Calibri", size=11, bold=True, color="000000")
    total_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    total_font = Font(name="Calibri", size=11, bold=True, color="000000")

    success_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    success_font = Font(name="Calibri", size=10, bold=True, color="375623")
    pending_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    pending_font = Font(name="Calibri", size=10, color="7F6000")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # ============================================================
    # SHEET 1: REKAPITULASI KEJURUAN
    # ============================================================
    ws_rekap = wb.active
    ws_rekap.title = "Rekapitulasi Kejuruan"
    ws_rekap.views.sheetView[0].showGridLines = True

    # Judul Dokumen
    ws_rekap.merge_cells("A1:G1")
    ws_rekap["A1"] = "BALAI PELATIHAN VOKASI DAN PRODUKTIVITAS (BPVP) KENDARI"
    ws_rekap["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F4E78")
    ws_rekap["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws_rekap.merge_cells("A2:G2")
    ws_rekap["A2"] = f"LAPORAN REKAPITULASI CAPAIAN KUOTA PELATIHAN VOKASI — DICETAK: {datetime.now().strftime('%d %B %Y %H:%M')}"
    ws_rekap["A2"].font = Font(name="Calibri", size=10, italic=True, color="595959")
    ws_rekap["A2"].alignment = Alignment(horizontal="center", vertical="center")

    rekap_headers = [
        "NO",
        "KEJURUAN / BIDANG",
        "TOTAL PAKET",
        "TARGET KUOTA (SISWA)",
        "SISWA TERDAFTAR",
        "SISA KUOTA",
        "PERSENTASE CAPAIAN",
    ]

    for col_idx, h_text in enumerate(rekap_headers, start=1):
        cell = ws_rekap.cell(row=4, column=col_idx, value=h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Agregasi data per kejuruan
    if not df.empty:
        summary_df = (
            df.groupby("BIDANG")
            .agg(
                paket=("PROGRAM PELATIHAN", "count"),
                terisi=("PESERTA", "sum"),
            )
            .reset_index()
            .sort_values(by="paket", ascending=False)
        )
        summary_df["target"] = summary_df["paket"] * kapasitas_default
        summary_df["sisa"] = summary_df.apply(lambda r: max(0, r["target"] - r["terisi"]), axis=1)
        summary_df["persen"] = summary_df.apply(
            lambda r: (r["terisi"] / r["target"] * 100) if r["target"] > 0 else 0.0, axis=1
        )

        curr_row = 5
        for idx, r in summary_df.iterrows():
            ws_rekap.cell(row=curr_row, column=1, value=idx + 1).alignment = Alignment(horizontal="center")
            ws_rekap.cell(row=curr_row, column=2, value=str(r["BIDANG"]))
            ws_rekap.cell(row=curr_row, column=3, value=int(r["paket"])).number_format = "#,##0"
            ws_rekap.cell(row=curr_row, column=4, value=int(r["target"])).number_format = "#,##0"
            ws_rekap.cell(row=curr_row, column=5, value=int(r["terisi"])).number_format = "#,##0"
            ws_rekap.cell(row=curr_row, column=6, value=int(r["sisa"])).number_format = "#,##0"
            
            pct_cell = ws_rekap.cell(row=curr_row, column=7, value=r["persen"] / 100.0)
            pct_cell.number_format = "0.0%"
            pct_cell.alignment = Alignment(horizontal="right")

            for c in range(1, 8):
                ws_rekap.cell(row=curr_row, column=c).border = thin_border
            curr_row += 1

        # Baris Total Keseluruhan
        ws_rekap.cell(row=curr_row, column=1, value="")
        ws_rekap.cell(row=curr_row, column=2, value="TOTAL KESELURUHAN")
        ws_rekap.cell(row=curr_row, column=3, value=f"=SUM(C5:C{curr_row-1})")
        ws_rekap.cell(row=curr_row, column=4, value=f"=SUM(D5:D{curr_row-1})")
        ws_rekap.cell(row=curr_row, column=5, value=f"=SUM(E5:E{curr_row-1})")
        ws_rekap.cell(row=curr_row, column=6, value=f"=SUM(F5:F{curr_row-1})")
        total_pct = ws_rekap.cell(row=curr_row, column=7, value=f"=E{curr_row}/D{curr_row}")
        total_pct.number_format = "0.0%"

        for c in range(1, 8):
            cell = ws_rekap.cell(row=curr_row, column=c)
            cell.fill = total_fill
            cell.font = total_font
            cell.border = thin_border

    # ============================================================
    # SHEET 2: DATA DETAIL PELATIHAN
    # ============================================================
    ws_data = wb.create_sheet(title="Data Pelatihan")
    ws_data.views.sheetView[0].showGridLines = True

    export_cols = [
        "NO",
        "BIDANG",
        "PROGRAM PELATIHAN",
        "STATUS SELESAI",
        "PESERTA",
        "LOKASI",
        "TANGGAL MULAI",
        "TANGGAL SELESAI",
        "RENCANA MULAI",
        "RENCANA SELESAI",
        "KETERANGAN",
    ]

    for col_idx, h_text in enumerate(export_cols, start=1):
        cell = ws_data.cell(row=1, column=col_idx, value=h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    curr_row = 2
    for _, row in df.iterrows():
        is_done = bool(row.get("CHECK", False))
        no_val = ""
        try:
            if not pd.isna(row.get("NO")):
                no_val = int(float(row.get("NO")))
        except Exception:
            no_val = str(row.get("NO", ""))

        ws_data.cell(row=curr_row, column=1, value=no_val).alignment = Alignment(horizontal="center")
        ws_data.cell(row=curr_row, column=2, value=str(row.get("BIDANG", "")))
        ws_data.cell(row=curr_row, column=3, value=str(row.get("PROGRAM PELATIHAN", "")))

        status_cell = ws_data.cell(row=curr_row, column=4, value="SELESAI" if is_done else "BELUM SELESAI")
        status_cell.alignment = Alignment(horizontal="center")
        if is_done:
            status_cell.fill = success_fill
            status_cell.font = success_font
        else:
            status_cell.fill = pending_fill
            status_cell.font = pending_font

        pst_val = 0
        try:
            pst_val = int(row.get("PESERTA", 0))
        except Exception:
            pst_val = 0
        ws_data.cell(row=curr_row, column=5, value=pst_val).number_format = "#,##0"
        ws_data.cell(row=curr_row, column=6, value=str(row.get("LOKASI", "")))
        ws_data.cell(row=curr_row, column=7, value=str(row.get("TANGGAL MULAI", "")))
        ws_data.cell(row=curr_row, column=8, value=str(row.get("TANGGAL SELESAI", "")))
        ws_data.cell(row=curr_row, column=9, value=str(row.get("RENCANA MULAI", "")))
        ws_data.cell(row=curr_row, column=10, value=str(row.get("RENCANA SELESAI", "")))
        ws_data.cell(row=curr_row, column=11, value=str(row.get("KETERANGAN", "")))

        for c in range(1, 12):
            if c != 4:  # status cell sudah diformat warnanya
                ws_data.cell(row=curr_row, column=c).border = thin_border
            else:
                status_cell.border = thin_border
        curr_row += 1

    # Atur lebar kolom otomatis untuk kedua sheet
    for ws in [ws_rekap, ws_data]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if cell.row in [1, 2] and ws == ws_rekap:
                    continue  # skip baris judul merged
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
